"""20 分鐘與老師討論用的提示卡：同一份內容輸出成網頁（HTML）與 Word。

內容是「提示」而不是逐字稿：每段給時間、要打開的圖、要指的程式行、
可以直接說出口的英文關鍵句（附中文意思）、概念錨點、要引用的數字與銜接句。
程式節錄依行號從原始檔讀取；原始碼改動後要重新檢查 SEGMENTS 內的行號。

執行：.venv/bin/python presentations/build_discussion_guide.py
輸出：outputs/discussion_guide/

Cue cards for a 20-minute discussion with the instructor: the same content as a web page (HTML) and Word.

The content is cues, not a script: each section gives the time, the figure to open, the code lines to point at,
English key sentences to say (with the Chinese meaning), concept anchors, numbers to cite and transitions.
Code excerpts are read from the source by line number; recheck the line numbers in SEGMENTS after changing the source.

Run: .venv/bin/python presentations/build_discussion_guide.py
Output: outputs/discussion_guide/
"""

from __future__ import annotations

import html
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs" / "discussion_guide"
FIG_DIR = OUT / "figs"
WEB_WIDTH = 1800  # 網頁與 Word 共用的縮圖寬度（px） / thumbnail width shared by the web page and Word (px)

FIGURES = {
    "single_cell_rules": (
        "outputs/mathematical_model/single_cell_rules.png",
        "圖 06 單細胞規則",
    ),
    "tumor_dashboard": (
        "outputs/tumor_dashboard.png",
        "圖 01 治療下的腫瘤（30 次試驗）",
    ),
    "muller": (
        "outputs/genotype_evolution/muller_top_genotypes_en.png",
        "Muller 圖 基因型組成隨時間",
    ),
    "matrices_to_flow": (
        "outputs/mathematical_model/matrices_to_flow.png",
        "圖 07 期望矩陣 M → 向量場",
    ),
    "phase_field_burden": (
        "outputs/phase_plots/phase_field_burden_logit_en.png",
        "相圖 05 負荷 × 抗藥比例：理論 vs 實際",
    ),
    "phase_field_fitness": (
        "outputs/phase_plots/phase_field_fitness_logit_en.png",
        "相圖 06 適應度 × 抗藥比例：演化與篩選的分工",
    ),
    "relapse_prediction": (
        "outputs/therapy_analysis/relapse_time_prediction.png",
        "圖 09 復發時間預測",
    ),
    "goldie_coldman": (
        "outputs/therapy_analysis/goldie_coldman.png",
        "圖 08 Goldie–Coldman 治癒曲線",
    ),
    "sensitivity": (
        "outputs/therapy_analysis/parameter_sensitivity.png",
        "圖 10 k_R × μ 敏感度",
    ),
    "one_hit_vs_two_hit": (
        "outputs/one_hit_vs_two_hit.png",
        "圖 03 One-hit vs two-hit",
    ),
    "death_rate": ("outputs/death_rate_heterogeneity.png", "圖 02 死亡率與異質性"),
    "adaptive_therapy": (
        "outputs/therapy_analysis/adaptive_therapy.png",
        "圖 11 適應性治療",
    ),
}

THESIS = (
    "Therapy does not create resistance — it selects resistance that is already there.",
    "治療不製造抗藥性，而是篩選本來就存在的抗藥細胞。",
)

# 每段：時間、標題、目標，以及數個「節拍」（beat）。每個 beat 對應一張圖或一段程式。
# Each section: time, title, goal and several beats. Each beat matches one figure or one code excerpt.
# says = (可直接說的英文句, 中文意思)；point = 指圖的順序；code = (檔案, 起, 迄, 強調行)。
# says = (English sentence to say, Chinese meaning); point = pointing order; code = (file, start, end, highlighted lines).
SEGMENTS = [
    dict(
        id="open",
        start=0.0,
        end=1.5,
        title="開場：一句話主張",
        goal="先給結論與模型範圍，讓老師知道接下來每張圖都在支持同一個主張。",
        beats=[
            dict(
                fig=None,
                says=[
                    THESIS,
                    (
                        (
                            "I built a discrete-generation branching process: each cell"
                            " dies, or makes two or three daughters."
                        ),
                        "離散世代分枝過程：每個細胞死亡，或產生 2／3 個子細胞。",
                    ),
                    (
                        (
                            "Twenty genes: three oncogenes, two tumor suppressors with"
                            " two alleles each, one resistance gene, fourteen"
                            " passengers."
                        ),
                        (
                            "20 個基因：3 oncogene、2 TSG（各兩個等位基因）、1"
                            " 抗藥基因、14 passenger。"
                        ),
                    ),
                    (
                        (
                            "Parameters are teaching assumptions, not clinical fits — I"
                            " care about mechanisms and trends."
                        ),
                        "參數是教學假設，重點是機制與趨勢，不是臨床預測。",
                    ),
                ],
                concepts=[
                    "branching process",
                    "selection vs. induction",
                    "教學假設 ≠ 校準",
                ],
                numbers=["初始 100 細胞", "20,000 細胞開始治療", "每組 30 次試驗"],
            )
        ],
        transition=(
            "Let me start with the single rule every result is built from.",
            "先從所有結果共用的那一條單細胞規則開始。",
        ),
    ),
    dict(
        id="model",
        start=1.5,
        end=5.5,
        title="模型與程式：一個細胞、一代",
        goal=(
            "用圖 06 講規則，用三段程式說明規則怎麼變成模擬；強調設計選擇（計數字典、二項抽樣）。"
        ),
        beats=[
            dict(
                fig="single_cell_rules",
                point=[
                    "A：一個父細胞 → 0／2／3 子細胞，機率 0.300／0.672／0.028",
                    "A 下方：m = f′(1) = 1.428 > 1 → 平均會長大",
                    "s = f(s) ⇒ s = 0.421：單一細胞家系仍有 42% 會滅絕",
                    "C 表：oncogene、TSG 讓 m 由 1.43 → 1.68 → 1.96；治療欄 m_ON",
                ],
                says=[
                    (
                        (
                            "Everything starts from one offspring law: f(z) = 0.3 +"
                            " 0.672 z² + 0.028 z³."
                        ),
                        "所有結果都由同一個子代分布 f(z) 推出。",
                    ),
                    (
                        (
                            "The mean m = 1.428 tells growth; the fixed point of f"
                            " tells extinction."
                        ),
                        "平均 m 決定成長；f 的不動點決定滅絕機率。",
                    ),
                    (
                        (
                            "Look at the last row: a resistant cell is neutral without"
                            " therapy, m = 1.428, but keeps 1.371 on therapy."
                        ),
                        "抗藥細胞沒治療時是中性（m 相同），治療下仍 m = 1.371 > 1。",
                    ),
                    (
                        (
                            "A sensitive cell on therapy has m = 0.14, so its lineage"
                            " dies with probability one."
                        ),
                        "敏感細胞治療下 m = 0.14，家系必然滅絕。",
                    ),
                ],
                concepts=[
                    "probability generating function (PGF)",
                    "mean offspring m：> 1 長、< 1 縮",
                    "extinction probability = smallest root of s = f(s)",
                    "resistance is neutral off-therapy",
                ],
                numbers=[
                    "m_OFF = 1.428",
                    "m_ON(敏感) = 0.143",
                    "m_ON(抗藥) = 1.371",
                    "家系滅絕 0.421／治療下抗藥 0.479",
                ],
            ),
            dict(
                fig=None,
                code=(
                    "simulation.py",
                    148,
                    165,
                    [151, 152, 153, 154, 155, 160, 161, 162, 163],
                ),
                code_title="基因型 → 表型：oncogene 與 TSG 各自降低死亡率、提高三子機率",
                says=[
                    (
                        (
                            "Genotype maps to just two numbers, death d and"
                            " triple-division probability q."
                        ),
                        "基因型最後只變成兩個數字：死亡率 d 與三子細胞機率 q。",
                    ),
                    (
                        "Each class of driver counts once, and the effects add.",
                        "每一類 driver 只算一次，兩類效果可相加。",
                    ),
                ],
                concepts=["genotype → phenotype map", "additive fitness effects"],
                numbers=[
                    "d：0.30 → 0.22（oncogene）→ 0.14（+TSG）",
                    "q：0.04 → 0.16 → 0.28",
                ],
            ),
            dict(
                fig=None,
                code=(
                    "simulation.py",
                    306,
                    326,
                    [312, 313, 314, 315, 316, 317, 319, 320, 322, 326],
                ),
                code_title="一代更新：治療殺傷 → 二項抽樣存活與三子 → 子代突變",
                says=[
                    (
                        (
                            "I store the tumor as genotype → count, not as individual"
                            " cells."
                        ),
                        "腫瘤以「基因型 → 細胞數」儲存，不逐一追蹤細胞。",
                    ),
                    (
                        (
                            "Two binomial draws per genotype are exactly equivalent to"
                            " flipping a coin for every cell — just faster."
                        ),
                        "每個基因型抽兩次二項分布，與逐細胞擲硬幣等價，只是更快。",
                    ),
                    (
                        (
                            "Therapy acts on natural survivors, so total death is 1 −"
                            " (1−d)(1−k)."
                        ),
                        "治療作用在自然存活者上，總死亡率 1 − (1−d)(1−k)。",
                    ),
                    (
                        (
                            "Mutations happen in daughters after division; their effect"
                            " starts next generation."
                        ),
                        "突變發生在分裂後的子細胞，效果下一代才生效。",
                    ),
                ],
                concepts=[
                    "binomial thinning",
                    "synchronous generations",
                    "timing: select → divide → mutate",
                ],
                numbers=["殺傷 k：敏感 0.90、抗藥 0.04"],
            ),
            dict(
                fig=None,
                code=("simulation.py", 188, 199, [193, 196, 197]),
                code_title="突變：逐位元二項抽樣，以 bitmask OR 設定新突變",
                says=[
                    (
                        "A genotype is a 22-bit integer; mutation is a bitwise OR.",
                        "基因型是 22 位元整數（20 基因 + 2 個 TSG 第二等位基因）；突變就是 OR。",
                    ),
                    (
                        (
                            "Mutation only relabels cells — the total count is"
                            " conserved, and that is unit-tested."
                        ),
                        "突變只重新分類細胞，總數守恆，並有單元測試。",
                    ),
                ],
                concepts=[
                    "bitmask genotype",
                    "irreversible mutations",
                    "μ = 10⁻³ per locus per daughter",
                ],
                numbers=["22 bits = 20 loci + 2 extra TSG alleles"],
            ),
        ],
        transition=(
            (
                "With that rule, here is what continuous therapy does across thirty"
                " independent tumors."
            ),
            "有了規則，來看 30 個獨立腫瘤在持續治療下的結果。",
        ),
    ),
    dict(
        id="therapy",
        start=5.5,
        end=9.5,
        title="核心結果：治療、谷底、復發",
        goal="用圖 01 + Muller 圖證明：抗藥細胞治療前就存在，治療只是讓它們接手。",
        beats=[
            dict(
                fig="tumor_dashboard",
                point=[
                    "A：藍線（敏感）在第 14–16 代治療後崩落；橘線（抗藥）幾乎不受影響",
                    "A：黑線總量 → 谷底 → 回升 = 復發",
                    "C：治療開始時抗藥比例中位數 1.2%，30/30 試驗 > 0",
                    "B：有效基因型數治療後先下降（瓶頸）再上升",
                ],
                says=[
                    (
                        (
                            "In all thirty trials, resistant cells were already present"
                            " when therapy started — median 1.2 percent."
                        ),
                        "30 次試驗全部在治療前就有抗藥細胞，中位數 1.2%。",
                    ),
                    (
                        (
                            "Therapy removes the sensitive majority within about six"
                            " generations; the resistant minority keeps growing at 1.37"
                            " per generation."
                        ),
                        "敏感多數約六代內被清除；抗藥少數以每代 ×1.37 持續成長。",
                    ),
                    (
                        (
                            "So the nadir is not a cure — it is the moment the"
                            " resistant clone becomes the majority."
                        ),
                        "谷底不是治癒，而是抗藥族群變成多數的時刻。",
                    ),
                    (
                        (
                            "Diversity drops sharply — a selective bottleneck — then"
                            " rebuilds on the resistant background."
                        ),
                        "多樣性急降（選擇瓶頸），之後在抗藥背景上重新累積。",
                    ),
                ],
                concepts=[
                    "pre-existing resistance",
                    "selective sweep / bottleneck",
                    "nadir → relapse",
                ],
                numbers=[
                    "治療開始：第 14–16 代",
                    "谷底中位數 ≈ 814 細胞（416–2,003）",
                    "30/30 復發；中位數治療後 13 代（9–16）",
                ],
            ),
            dict(
                fig=None,
                code=("simulation.py", 341, 350, [347, 348, 349]),
                code_title="復發定義：先低於治療起點負荷，再回到該負荷",
                says=[
                    (
                        (
                            "Relapse is defined operationally: first drop below the"
                            " starting burden, then return to it."
                        ),
                        "復發是操作型定義：先降到起始負荷以下，再回到起始負荷。",
                    ),
                    (
                        (
                            "No relapse within the window is reported separately — it"
                            " is not called a cure."
                        ),
                        "觀察期內未復發另外標記，不等同治癒。",
                    ),
                ],
                concepts=["operational definition", "censoring"],
                numbers=[],
            ),
            dict(
                fig="muller",
                point=[
                    "上方：負荷曲線，標出治療開始（第 15 代）、谷底、復發（第 28 代）",
                    "治療前：灰色 WT 為主，珊瑚紅 ONC1 擴張；斜線（抗藥）色塊很細",
                    (
                        "治療後：斜線的抗藥基因型取代一切；灰色斜線 RES 內，"
                        "暖色系的 ONC1／ONC2／ONC3+RES 逐漸變大（oncogene 暖、TSG 綠、兩者紫）"
                    ),
                    "下排 6 個試驗：同樣的取代，但接手的基因型不同",
                ],
                says=[
                    (
                        (
                            "The winners after therapy are resistant genotypes that"
                            " also carry an oncogene."
                        ),
                        "治療後的贏家是同時帶 oncogene 的抗藥基因型。",
                    ),
                    (
                        (
                            "Same outcome in every trial, different genotypes each time"
                            " — convergent phenotype, divergent genotype."
                        ),
                        "每次結果相同、接手的基因型不同：表型趨同、基因型分歧。",
                    ),
                    (
                        (
                            "Passengers hitchhike along with whichever resistant clone"
                            " happens to win."
                        ),
                        "Passenger 突變搭便車，跟著勝出的抗藥 clone 一起擴張。",
                    ),
                ],
                concepts=[
                    "clonal replacement",
                    "hitchhiking passengers",
                    "genotype containment ≠ lineage",
                ],
                numbers=["代表試驗 rep 4：治療 15 代、復發 28 代"],
            ),
        ],
        transition=(
            (
                "That was descriptive. Next I tried to make it predictive — with theory"
                " that has no fitted parameters."
            ),
            "以上是描述；接著用不擬合參數的理論讓它變成可預測。",
        ),
    ),
    dict(
        id="theory",
        start=9.5,
        end=14.0,
        title="從描述到預測：理論對照模擬",
        goal=(
            "用 M 矩陣、相圖 05／06、復發預測與 Goldie–Coldman"
            " 說明結果可由第一原理推出；這段最能展現深度。"
        ),
        beats=[
            dict(
                fig="matrices_to_flow",
                point=[
                    "左：Q（突變機率，列和 = 1）與 M = m·Q（期望子代，列和 = m）",
                    "B 無治療：箭頭全部往右上，S 與 R 一起長",
                    "C 持續治療：箭頭往左（S 消失）再往上（R 成長）",
                    "橘線 = 實際試驗，沿著箭頭走",
                ],
                says=[
                    (
                        (
                            "M is not a probability matrix; its rows sum to the mean"
                            " offspring m."
                        ),
                        "M 不是機率矩陣；列和是平均子代數 m。",
                    ),
                    (
                        (
                            "On therapy: ΔS = −0.857 S and ΔR = 0.371 R — the arrows"
                            " are just these two lines."
                        ),
                        "治療下 ΔS = −0.857S、ΔR = 0.371R，箭頭就是這兩條式子。",
                    ),
                    (
                        (
                            "The stochastic trajectories follow the expected flow — the"
                            " noise matters mostly when numbers are small."
                        ),
                        "隨機軌跡順著期望流場走；隨機性主要在數量很小時才重要。",
                    ),
                ],
                concepts=[
                    "P vs Q vs M",
                    "expected one-step drift",
                    "deterministic skeleton + noise",
                ],
                numbers=["ΔS = −0.8573 S", "ΔR = 0.0001428 S + 0.3709 R"],
            ),
            dict(
                fig="phase_field_burden",
                point=[
                    "上排 = 理論（藍）、下排 = 實際平均一步（橘，顏色 = 轉移數）；左無治療、右持續治療",
                    (
                        "x = log10(N+1)；y = 抗藥比例的 logit（刻度標成百分比）；虛線 S"
                        " = 0、點線 R = 0"
                    ),
                    "左：ρ 幾乎水平，抗藥中性，只有新突變讓它緩慢上升",
                    (
                        "右：灰色軌跡呈 C 形，先往左上（N 降、ρ 急升），再沿 S = 0 邊界往右上 = 復發"
                    ),
                ],
                says=[
                    (
                        (
                            "On logit coordinates the whole therapy response is one"
                            " C-shaped turn: burden falls while resistance takes over,"
                            " then burden climbs on an almost fully resistant tumor."
                        ),
                        (
                            "在 logit 座標下，治療反應是一個 C 形轉彎：負荷下降、"
                            "抗藥接手，之後幾乎全抗藥的腫瘤重新長大。"
                        ),
                    ),
                    (
                        (
                            "Theory says the resistant fraction rises about one decade"
                            " in logit per generation — log10 of w_R over w_S is 0.98."
                        ),
                        (
                            "理論：治療下 logit ρ 每代約上升一個 10 倍（log10 w_R/w_S ="
                            " 0.98）。"
                        ),
                    ),
                    (
                        (
                            "ρ → 1 is the stable direction in composition, but burden"
                            " has no fixed point; the one-million cap is computational,"
                            " not an equilibrium."
                        ),
                        "ρ → 1 是組成上的穩定方向，但負荷沒有不動點；一百萬上限是計算門檻，不是平衡。",
                    ),
                    (
                        (
                            "Where the theory and empirical rows differ, that is the"
                            " hidden genotype information and noise this 2-D projection"
                            " throws away."
                        ),
                        "上下排的差異來自被投影掉的基因型資訊與隨機性。",
                    ),
                ],
                concepts=[
                    "logit transform",
                    "2-D projection of a high-dimensional state",
                    "φ(E[X′]) vs E[φ(X′)]",
                ],
                numbers=[
                    "治療下 logit ρ +0.98／代",
                    "30 次試驗、1,089 次一步轉移",
                    "經驗箭頭：每格 ≥ 5 次轉移、≥ 2 次試驗",
                ],
            ),
            dict(
                fig="phase_field_fitness",
                point=[
                    "x 換成無治療平均適應度 w̄_off（oncogene 累積 → 往右），y 同相圖 05",
                    "左（無治療）：箭頭水平往右 = 演化；ρ 停在約 1%",
                    "右（持續治療）：箭頭幾乎垂直往上 = 篩選",
                    "右下實際：先垂直衝到 > 99.99%，之後才沿頂端往右",
                ],
                says=[
                    (
                        (
                            "This plot separates two processes: horizontal motion is"
                            " evolution — oncogenes accumulating; vertical motion is"
                            " selection for resistance."
                        ),
                        "這張圖把兩件事分開：水平移動是演化（oncogene 累積），垂直移動是抗藥篩選。",
                    ),
                    (
                        (
                            "Without therapy the tumor drifts right — fitness rises"
                            " while the resistant fraction stays near one percent."
                        ),
                        "無治療時腫瘤往右漂：適應度上升，抗藥比例維持約 1%。",
                    ),
                    (
                        (
                            "Therapy turns the flow vertical; only after the tumor is"
                            " almost fully resistant does fitness evolution resume."
                        ),
                        "治療讓流向變垂直；幾乎全抗藥之後，適應度演化才繼續。",
                    ),
                    (
                        (
                            "That is why relapsed tumors are both resistant and fitter"
                            " — the Muller plot showed resistant clones carrying an"
                            " oncogene winning."
                        ),
                        (
                            "所以復發腫瘤既抗藥又更適應，對應 Muller 圖中帶 oncogene"
                            " 的抗藥 clone 勝出。"
                        ),
                    ),
                ],
                concepts=[
                    "evolution vs selection",
                    "orthogonal axes",
                    "mean fitness w̄",
                ],
                numbers=[
                    "w̄_off ≈ 1.43 → 1.65",
                    "理論截面：N = 20,000、oncogene 比例與抗藥獨立",
                ],
            ),
            dict(
                fig="relapse_prediction",
                point=[
                    "A：灰點 = 只用抗藥細胞數的基準預測；紫點 = 依基因型組成的預測，貼近對角線",
                    "B：R₀ 越多 → 越早復發（Spearman −0.70）",
                    "C：成長率對抗藥比例 ρ；在 ρ* ≈ 70% 穿過零 = 谷底",
                ],
                says=[
                    (
                        (
                            "Knowing only the composition at therapy start, I can"
                            " predict the relapse generation."
                        ),
                        "只知道治療開始時的組成，就能預測復發世代。",
                    ),
                    (
                        (
                            "A single growth rate gives a mean error of two"
                            " generations; letting each resistant genotype grow at its"
                            " own rate cuts it to 0.2."
                        ),
                        "單一成長率誤差 2 代；依各基因型自身成長率，誤差降到 0.2 代。",
                    ),
                    (
                        (
                            "The nadir happens exactly when the resistant fraction"
                            " crosses ρ* = (1 − m_S)/(m_R − m_S), about seventy"
                            " percent."
                        ),
                        "抗藥比例越過 ρ* ≈ 70% 的那一代就是谷底。",
                    ),
                ],
                concepts=[
                    "mixture growth model",
                    "ρ* threshold",
                    "Spearman rank correlation",
                ],
                numbers=[
                    "MAE 2.0 → 0.2 代",
                    "30 次中 24 次預測到正確世代",
                    "oncogene 抗藥細胞 ×1.62/代",
                ],
            ),
            dict(
                fig="goldie_coldman",
                point=[
                    (
                        "A：x = 開始治療時的大小（對數），y = 治癒機率；黑色階梯 ="
                        " 理論，橘點 = 模擬（200 次、Wilson 95% CI）"
                    ),
                    "50% 治癒點 ≈ 416 細胞",
                    "B：已存在的抗藥細胞；平均 > 中位數 → Luria–Delbrück jackpot",
                ],
                says=[
                    (
                        (
                            "This is the Goldie–Coldman result: the probability of cure"
                            " falls steeply with tumor size at treatment."
                        ),
                        "Goldie–Coldman：治療時腫瘤越大，治癒機率急降。",
                    ),
                    (
                        (
                            "The theory has no fitted parameters — Poisson arrival of"
                            " mutant clones, each surviving with 1 − f⁽ˢ⁾(q_on)."
                        ),
                        (
                            "理論不擬合參數：突變 clone 以 Poisson 出現，各自以 1 −"
                            " f⁽ˢ⁾(q_on) 存活。"
                        ),
                    ),
                    (
                        (
                            "The mean number of resistant cells is driven by rare early"
                            " jackpots, exactly as in Luria–Delbrück."
                        ),
                        "抗藥細胞平均數由少數早期 jackpot 主導，正如 Luria–Delbrück。",
                    ),
                ],
                concepts=[
                    "Goldie–Coldman hypothesis",
                    "Poisson mutation process",
                    "Luria–Delbrück fluctuation",
                ],
                numbers=[
                    "100 細胞 → 99.5% 治癒",
                    "1,000 → 10%",
                    "5,000 → 0%",
                    "q_on = 0.479",
                ],
            ),
            dict(
                fig=None,
                code=("analysis.py", 197, 205, [201, 202]),
                code_title="家系滅絕機率：由 0 迭代 s ← f(s) 到最小不動點",
                says=[
                    (
                        (
                            "The theory reuses the same offspring law as the simulation"
                            " — iterate s ← f(s) from zero."
                        ),
                        "理論重用模擬的同一個子代分布：從 0 迭代 s ← f(s)。",
                    ),
                ],
                concepts=["fixed-point iteration", "independent validation"],
                numbers=[],
            ),
            dict(
                fig="sensitivity",
                optional=True,
                point=[
                    "A：x = 抗藥細胞殺傷率 k_R，y = 突變率 μ；顏色 = 治癒機率",
                    (
                        "實線 k = 0.30：單純抗藥細胞停止成長；虛線 k = 0.41：帶"
                        " oncogene 的抗藥細胞也停止"
                    ),
                ],
                says=[
                    (
                        (
                            "Stopping the average resistant cell is not enough; you"
                            " have to stop the fittest one."
                        ),
                        "讓一般抗藥細胞停止成長不夠，必須壓住最適應的那一種。",
                    ),
                ],
                concepts=["fittest-clone threshold", "(1 − k_R)·m = 1"],
                numbers=["k_R = 0.30 vs 治癒需要 ≈ 0.41"],
            ),
        ],
        transition=(
            (
                "Now the scenarios from the assignment: one-hit versus two-hit, and"
                " death rate versus heterogeneity."
            ),
            "接著是作業建議的情境：one-hit vs two-hit，以及死亡率與異質性。",
        ),
    ),
    dict(
        id="scenarios",
        start=14.0,
        end=17.0,
        title="作業情境：Two-hit、死亡率",
        goal="每張圖只講一個結論與一個機制；保留對照組，避免過度解讀。",
        beats=[
            dict(
                fig="one_hit_vs_two_hit",
                point=[
                    "A：三條生長曲線幾乎重疊，one-hit（橘）最上面",
                    "B：TSG 達標比例 one-hit 第 18 代 27.9%；two-hit 0.07%，與對照差不多",
                    "C：第 18 代負荷，one-hit 中位數最高",
                ],
                says=[
                    (
                        (
                            "With one hit, TSG loss spreads within a few generations;"
                            " with two hits it barely appears in eighteen."
                        ),
                        "One-hit 幾代內擴散；two-hit 18 代內幾乎沒出現。",
                    ),
                    (
                        (
                            "Two-hit needs both alleles of the same gene — roughly μ"
                            " squared unless the first hit expands first."
                        ),
                        "Two-hit 需要同一基因兩個等位基因，機率約 μ²，除非第一擊先擴張。",
                    ),
                    (
                        (
                            "That is Knudson's two-hit logic: recessive suppressors"
                            " delay the selective advantage."
                        ),
                        "這就是 Knudson two-hit：隱性抑癌基因延遲選擇優勢。",
                    ),
                ],
                concepts=[
                    "Knudson two-hit",
                    "haploinsufficiency（one-hit 的簡化）",
                    "no-advantage control",
                ],
                numbers=[
                    "第 18 代中位數：對照 76,446／one-hit 94,618／two-hit 73,879",
                    "TSG 達標 0.05%／27.9%／0.07%",
                ],
            ),
            dict(
                fig=None,
                code=("simulation.py", 131, 143, [134, 135, 141]),
                code_title="Two-hit：同一 TSG 的兩個等位基因位元都要失活",
                says=[
                    (
                        (
                            "The second allele lives in an extra bit beyond the twenty"
                            " loci, so two different genes with one hit each do not"
                            " count."
                        ),
                        "第二等位基因放在 20 個位點之外的額外位元；兩個不同基因各一擊不算。",
                    ),
                ],
                concepts=["allele-level bookkeeping"],
                numbers=[],
            ),
            dict(
                fig="death_rate",
                point=[
                    "A：同樣 5,000 細胞樣本，死亡率 0.20 → 0.45，Shannon H 1.06 → 2.83",
                    "B：長到 5,000 細胞的代數 8 → 14 → 26.5",
                    "x 軸下方：淨成長 ×1.63 → ×1.33 → ×1.12／代",
                ],
                says=[
                    (
                        (
                            "Higher death rate means more turnover: more divisions to"
                            " reach the same size, so more mutations."
                        ),
                        "死亡率高 = 周轉高：長到同樣大小需要更多分裂，累積更多突變。",
                    ),
                    (
                        (
                            "Same size, very different history — tumor size alone does"
                            " not tell you its diversity."
                        ),
                        "同樣大小、歷史不同；只看大小無法知道異質性。",
                    ),
                ],
                concepts=["cell turnover", "equal-size sampling", "Shannon diversity"],
                numbers=["H：1.06／1.80／2.83", "代數：8／14／26.5"],
            ),
        ],
        transition=(
            (
                "If resistance is selected, could we select less? That is the idea"
                " behind adaptive therapy."
            ),
            "如果抗藥性是被篩選出來的，能否少篩選一點？這就是適應性治療。",
        ),
    ),
    dict(
        id="adaptive",
        start=17.0,
        end=18.5,
        title="延伸：適應性治療",
        goal="說明一個反直覺結果：間歇治療只有在敏感細胞能壓制抗藥細胞時才有用。",
        beats=[
            dict(
                fig="adaptive_therapy",
                point=[
                    "A：持續治療；B：間歇治療（陰影 = 用藥中），敏感細胞被保留下來",
                    "C：四種設定的惡化時間；右側比值 ×0.85 → ×1.14",
                ],
                says=[
                    (
                        (
                            "Without competition, pausing therapy only gives resistant"
                            " cells a head start — adaptive is worse."
                        ),
                        "沒有競爭時，暫停治療只是讓抗藥細胞先跑，間歇治療反而較差。",
                    ),
                    (
                        (
                            "With a carrying capacity, sensitive cells hold resistant"
                            " ones back; adaptive delays progression with about half"
                            " the drug."
                        ),
                        "有承載量競爭時，敏感細胞壓制抗藥細胞；間歇治療延後惡化，藥量約減半。",
                    ),
                ],
                concepts=[
                    "competitive release",
                    "carrying capacity K",
                    "cost of resistance",
                ],
                numbers=[
                    "無競爭 ×0.85",
                    "K = 30,000：14 → 16 代（×1.14）",
                    "加抗藥代價：17.5 → 20 代",
                ],
            ),
            dict(
                fig=None,
                code=("simulation.py", 168, 173, [172, 173]),
                code_title="承載量競爭：額外死亡 c·N/K，使無突變細胞在 N = K 時 m = 1",
                says=[
                    (
                        (
                            "Competition is one extra death term, scaled so a"
                            " mutation-free cell exactly replaces itself at N = K."
                        ),
                        "競爭只是一項額外死亡率，讓無突變細胞在 N = K 時剛好自我替換。",
                    ),
                ],
                concepts=["density dependence"],
                numbers=["c ≈ 0.21"],
            ),
        ],
        transition=(
            "Let me close with what the model cannot say.",
            "最後說明模型的限制。",
        ),
    ),
    dict(
        id="close",
        start=18.5,
        end=20.0,
        title="限制與收尾",
        goal="主動講限制（建立可信度），用三個 take-home 收尾，邀請老師提問。",
        beats=[
            dict(
                fig=None,
                says=[
                    (
                        (
                            "Limitations: no space, no immune system, no quiescent"
                            " cells, and the mutation rate is inflated for a short"
                            " demo."
                        ),
                        "限制：無空間、無免疫、無靜止細胞；突變率為短時間演示而放大。",
                    ),
                    (
                        (
                            "Bands are ten-to-ninety percent ranges across trials, not"
                            " confidence intervals, and I did not pick seeds."
                        ),
                        "陰影是 10–90% 試驗間範圍，不是信賴區間；沒有挑選 seed。",
                    ),
                    (
                        (
                            "A genotype in my model is not a lineage — the same"
                            " genotype can arise independently."
                        ),
                        "模型中的基因型不等於家系，同一基因型可獨立出現多次。",
                    ),
                    (
                        (
                            "Three take-homes: resistance is pre-existing; its amount"
                            " at treatment predicts relapse; size and turnover set the"
                            " odds of cure."
                        ),
                        "三個重點：抗藥性預先存在；治療時的量可預測復發；大小與周轉決定治癒機會。",
                    ),
                ],
                concepts=[
                    "model scope",
                    "variability ≠ uncertainty",
                    "next steps: two-drug therapy (μ²), space, quiescence",
                ],
                numbers=["1,000,000 細胞是計算上限，不是承載量"],
            ),
        ],
        transition=(
            "I'm happy to run any scenario live if you'd like to see it.",
            "如果您想看，我可以現場跑任何情境。",
        ),
    ),
]

QA = [
    (
        "Why sample by genotype counts instead of simulating each cell?",
        (
            "每個細胞獨立 → 同基因型的存活數就是 Binomial(n, 1−d)；與逐細胞完全等價，"
            "速度快很多（simulation.py:319–322）。"
        ),
    ),
    (
        "Isn't μ = 10⁻³ per gene far too high?",
        (
            "是。真實約 10⁻⁶／基因／分裂；放大是為了在 ~30 代內看到現象。圖 10 顯示 μ"
            " 越高治癒越難；趨勢不變，時間尺度會變。"
        ),
    ),
    (
        "Why can a cell only make 0, 2 or 3 daughters — no 1?",
        "照作業設定（分裂 ≥ 2）且父細胞被取代；所以沒有靜止或休眠細胞，這也是限制之一。",
    ),
    (
        "Does therapy cause resistance mutations in your model?",
        (
            "不會。突變率不受治療影響；抗藥在治療前就以 Poisson 出現（圖 01-C：30/30 >"
            " 0；圖 08-B）。"
        ),
    ),
    (
        "Why does relapse timing vary between trials?",
        (
            "R₀ 的 jackpot 分布（Luria–Delbrück）以及抗藥細胞是否帶 oncogene（圖"
            " 09-B：Spearman −0.70、−0.51）。"
        ),
    ),
    (
        "How do you know the code is correct?",
        (
            "單元測試（突變守恆、等位基因、可重現性）；理論與模擬吻合（圖 08）；1,604"
            " 個網格點的期望值與完整模型核對。"
        ),
    ),
    (
        "What is the difference between P and M?",
        "P 是細胞數的機率轉移矩陣（列和 1）；M 是期望子代矩陣（列和 = m），不是機率。",
    ),
    (
        "Would combination therapy help?",
        (
            "需同時對兩藥抗藥 → 約 μ² 的事件；預期讓 50% 治癒點大幅右移。"
            "可加第二個抗藥基因驗證，是下一步。"
        ),
    ),
    (
        "Why compare diversity at a fixed size instead of fixed time?",
        "固定時間會混入大小差異；固定 5,000 細胞樣本只留下周轉的效果（圖 02）。",
    ),
]

CHEAT = [
    ("m（無突變，無治療）", "1.428"),
    ("家系滅絕機率", "0.421"),
    ("治療下 m：敏感／抗藥", "0.143／1.371"),
    ("治療開始時抗藥比例", "中位數 1.2%，30/30 > 0"),
    ("谷底中位數", "≈ 814 細胞（416–2,003）"),
    ("復發延遲", "中位數 13 代（9–16），30/30 復發"),
    ("50% 治癒的治療起點", "≈ 416 細胞"),
    ("復發預測誤差", "2.0 → 0.2 代"),
    ("谷底比例 ρ*", "≈ 70%"),
    ("第 18 代負荷：對照／one／two", "76,446／94,618／73,879"),
    ("Shannon H（d = 0.20／0.35／0.45）", "1.06／1.80／2.83"),
    ("適應性治療（K = 30,000）", "×1.14，約一半藥量"),
]

# 程式導覽：以類別卡（UML 風格）介紹架構。Config／Result 是真正的類別（dataclass）；
# Code tour: introduce the architecture with UML-style class cards. Config / Result are real classes (dataclasses);
# 其他檔案是無狀態的模組，這裡把「一個模組」畫成「一張類別卡」：職責 + 對外方法。
# the other files are stateless modules, drawn here as one class card per module: responsibility + public methods.
ARCH_INTRO = (
    "Config 與 Result 是真正的類別（dataclass），代表「參數」與「一次試驗的結果」兩個物件；"
    "其他檔案是無狀態的模組，這裡把每個模組畫成一張類別卡：上半是它保存的資料（屬性），"
    "下半是對外的方法與回傳值。整條資料流只有一個方向：Config → simulate() → Result → 彙整／理論 → 圖與 CSV。"
)

# 分層：(層名, 模組, 一句話職責)；網頁畫成由左到右的流程帶。
# Layers: (layer name, modules, one-line responsibility); drawn on the page as a left-to-right band.
ARCH_LAYERS = [
    ("入口", "main.py", "讀命令列參數，依序呼叫 5 個步驟"),
    ("實驗", "experiments.py", "用不同設定重複跑模型、彙整、讀寫 CSV"),
    ("模型", "simulation.py", "Config → simulate() → Result"),
    ("理論", "analysis.py", "由規則直接推出的公式與統計量"),
    ("呈現", "plot_*.py", "把彙整後的資料畫成圖 01–11"),
]

# 每張類別卡：kind（UML 刻板印象）、name、file、role、attrs [(名稱: 型別, 說明)]、
# Each class card: kind (UML stereotype), name, file, role, attrs [(name: type, description)],
# methods [(簽名 → 回傳, 說明)]。只列對外會用到的方法，內部小工具省略。
# methods [(signature → return, description)]. Only public methods; small internal helpers are omitted.
ARCH_CLASSES = [
    dict(
        kind="«module» 入口",
        name="main",
        file="main.py",
        role="唯一的執行入口；本身不含模型邏輯，只負責依序呼叫各步驟。",
        attrs=[
            ("STEPS", "results, model, phase, genotype, therapy"),
            (
                "ANALYSES",
                "goldie_coldman, relapse_prediction, sensitivity, adaptive（--analyses 可只跑部分）",
            ),
        ],
        methods=[
            (
                "main() / parse_args()",
                "讀 --only、--seed、--replicates…，決定要跑哪些步驟",
            ),
            (
                "run_results(args)",
                "主實驗 + 對照實驗 → 圖 01–03 與 CSV（其他步驟都讀它的 CSV）",
            ),
            ("run_model() / run_phase(args)", "數學圖 06／07 與相圖 01／05／06"),
            ("run_genotype() / run_therapy(args)", "基因型演化圖與治療延伸分析 08–11"),
            ("reproduced_trials()（快取）", "依 seed 重跑一次，後面三個步驟共用"),
        ],
    ),
    dict(
        kind="«dataclass, frozen»",
        name="Config",
        file="simulation.py",
        role="模型的全部參數；建立後不可修改（frozen），同一組參數可安全地分給 30 次試驗。",
        attrs=[
            ("n_genes: int = 20", "基因座數；另加 2 個位元給 TSG 的第二等位基因"),
            ("driver_genes = (0, 1, 2)", "3 個 oncogene；任一突變即活化"),
            (
                "suppressor_genes = (4, 5)",
                "2 個 TSG；suppressor_hits = 2 表示需兩個等位基因都失活",
            ),
            ("resistance_gene = 3", "抗藥基因；沒有治療時是中性"),
            (
                "death_rate = 0.30, triple_probability = 0.04",
                "基礎死亡率 d 與三子細胞機率 q",
            ),
            ("mutation_rate = 0.001", "每個子細胞、每個未突變位元、每代的突變機率 μ"),
            ("therapy_start_size = 20,000", "負荷首次達到此值就開始治療"),
            (
                "sensitive / resistant_therapy_kill = 0.90 / 0.04",
                "治療對敏感／抗藥細胞的額外殺傷率 k",
            ),
            (
                "initial_cells = 100, generations = 75, max_cells = 10⁶",
                "起始細胞數、觀察代數、計算上限",
            ),
        ],
        methods=[
            ("validate() → None", "檢查索引與機率是否合法；不合法就丟出 ValueError"),
        ],
    ),
    dict(
        kind="«dataclass»",
        name="Result",
        file="simulation.py",
        role="一次試驗的完整紀錄；由 simulate() 產生，之後所有彙整、理論與畫圖都讀它。",
        attrs=[
            ("config: Config", "產生這次試驗的參數（組合關係）"),
            ("seed: int", "亂數種子；同 seed 必得到同結果"),
            ("history: list[dict]", "每代摘要：total、resistant、sensitive、shannon…"),
            ("populations: list[dict[int, int]]", "每代的 {基因型: 細胞數} 快照"),
            ("therapy_generation: int | None", "開始治療的世代；None = 從未達門檻"),
            (
                "stop_reason: str",
                "extinction／cell_limit／generation_limit／target_reached",
            ),
            ("therapy_active: list[bool]", "每次轉移是否用藥（適應性治療會開關）"),
        ],
        methods=[],
    ),
    dict(
        kind="«module» 模型核心",
        name="simulation",
        file="simulation.py",
        role=(
            "細胞不是一個個物件，而是字典 {基因型: 細胞數}；基因型是整數位元遮罩，第 i 位元 = 1"
            " 代表基因 i 已突變。每代對每個基因型做二項抽樣，與逐細胞模擬等價但快很多。"
        ),
        attrs=[],
        methods=[
            (
                "simulate(config, seed, therapy=True) → Result",
                "主迴圈：一代一代推進，直到滅絕、達上限或時間到",
            ),
            (
                "growth_probabilities(genotype, config) → (d, q)",
                "基因型 → 表型：oncogene／TSG 降低 d、提高 q",
            ),
            (
                "mutate_offspring(offspring, config, rng) → dict",
                "子代逐位元以 μ 突變；細胞總數不變，只換基因型",
            ),
            (
                "summarize(population, generation, config) → dict",
                "算一代的負荷、抗藥數與多樣性，存進 history",
            ),
            (
                "therapy_outcome(result) → dict",
                "判讀復發：先降到治療前負荷以下，再回到該負荷",
            ),
            (
                "driver_count / suppressor_loss(genotype, config)",
                "位元運算判斷 oncogene 是否活化、TSG 是否失活",
            ),
        ],
    ),
    dict(
        kind="«module» 實驗管理",
        name="experiments",
        file="experiments.py",
        role="決定「跑哪些設定、跑幾次」，把多個 Result 整理成表格並寫檔；也負責從 CSV 讀回結果。",
        attrs=[("OUTPUTS: Path", "所有輸出的根目錄 outputs/")],
        methods=[
            (
                "run_therapy_trials(config, n, seed) → list[Result]",
                "主實驗：第 rep 次用 seed + 100000 + rep 呼叫 simulate",
            ),
            (
                "therapy_ensemble(trials) → (summary, raw)",
                "依世代彙整中位數與 10–90% 範圍（圖 01）",
            ),
            (
                "therapy_events(trials) → list[dict]",
                "每次試驗的治療起點、谷底、復發世代",
            ),
            (
                "run_tsg_comparison / run_death_rate_comparison",
                "對照實驗：圖 03 與圖 02",
            ),
            (
                "goldie_coldman / relapse_prediction / sensitivity / adaptive_experiment",
                "治療延伸分析 08／09／10／11",
            ),
            ("run_many(jobs, workers)", "多核心平行跑試驗（08、10、11 使用）"),
            (
                "load_trials / reproduce_trials(directory)",
                "讀回 CSV，或依記錄的 seed 重跑取得完整基因型",
            ),
            ("write_csv / write_json(path, …)", "寫出資料檔"),
        ],
    ),
    dict(
        kind="«module» 理論",
        name="analysis",
        file="analysis.py",
        role="不跑模擬，直接由規則算出理論值，用來和模擬結果互相核對。",
        attrs=[],
        methods=[
            ("offspring_law(d, q) → array", "單一細胞的子代分布 P(0), P(2), P(3)"),
            ("extinction(law) → float", "PGF 不動點 s = f(s)：家系滅絕機率"),
            (
                "fitness(genotype, config, therapy) → float",
                "期望子代數 m = (1−d)(1−k)(2+q)",
            ),
            ("markov_matrices(config) → dict", "圖 06／07 的矩陣 P、Q、M"),
            ("expected_next / compute_phase_field(config, …)", "一步期望 → 相圖向量場"),
            (
                "goldie_coldman_theory(config, size) → dict",
                "Poisson 近似的治癒機率（圖 08 理論線）",
            ),
            ("muller_layout / gene_frequency_series", "Muller 圖與基因頻率所需的整理"),
            ("wilson / spearman", "信賴區間與等級相關"),
        ],
    ),
    dict(
        kind="«module» × 6 呈現",
        name="plot_*",
        file="plot_style.py + 5 個繪圖檔",
        role="只畫圖，不做計算；輸入是 experiments／analysis 整理好的資料，輸出 PNG 與 SVG。",
        attrs=[],
        methods=[
            (
                "plot_style.set_style(lang) / save_lang(fig, …)",
                "共用字型、配色與雙語存檔",
            ),
            (
                "plot_results.plot_therapy_ensemble / death_rate / hit_comparison",
                "圖 01、02、03",
            ),
            ("plot_model.single_cell_figure / matrices_to_flow_figure", "圖 06、07"),
            ("plot_phase.plot_phase_portrait / plot_pair", "相圖 01、05、06"),
            (
                "plot_genotype.plot_muller / plot_gene_frequencies",
                "Muller 圖與基因頻率",
            ),
            (
                "plot_therapy.plot_goldie_coldman / relapse / sensitivity / adaptive",
                "圖 08–11",
            ),
        ],
    ),
]

# 執行流程：(呼叫, 說明)。第 4 步展開 simulate() 內一代發生的事。
# Execution flow: (call, description). Step 4 expands what happens in one generation inside simulate().
ARCH_FLOW = [
    ("$ .venv/bin/python main.py", "main() → parse_args() 決定要跑的步驟"),
    ("run_results(args)", "建立參數物件 config = Config(suppressor_hits=2)"),
    ("exp.run_therapy_trials(config, 30, seed)", "對 30 個不同 seed 呼叫 simulate()"),
    (
        "simulate(config, seed) 每一代",
        "summarize() 記錄 → 檢查是否停止 → 負荷 ≥ 20,000 就開治療 → 每個基因型："
        "growth_probabilities() 取 d、q → Binomial 抽存活數 → Binomial 抽三子數 →"
        " 全部分裂完才 mutate_offspring()",
    ),
    ("return Result", "history、populations、therapy_generation、stop_reason"),
    (
        "exp.therapy_ensemble / therapy_events",
        "多個 Result → 逐代摘要與每次試驗的事件 → write_csv",
    ),
    ("plot_results.plot_*", "畫圖 01–03"),
    (
        "run_model / run_phase / run_genotype / run_therapy",
        "讀回 CSV 或依 seed 重跑 → analysis 算理論值 → plot_* 畫圖 06–11",
    ),
]

ARCH_SAYS = [
    (
        "The code has one direction of data flow: a Config goes into simulate, a Result comes out, and everything else only reads Results.",
        "資料只往一個方向流：Config 進 simulate，Result 出來，其他程式都只讀 Result。",
    ),
    (
        "A tumor is a dictionary from genotype bitmask to cell count, so one generation is a few binomial draws per genotype.",
        "腫瘤是 {基因型位元遮罩: 細胞數}，所以一代只是每個基因型幾次二項抽樣。",
    ),
    (
        "Model, experiments, theory and plotting are separate files, so the theory can check the simulation independently.",
        "模型、實驗、理論、畫圖分檔，理論可以獨立核對模擬。",
    ),
]


# 將分鐘數格式化為分與秒的時間字串。
# Format minutes as a minutes:seconds string.
def fmt_time(minutes: float) -> str:
    return f"{int(minutes)}:{int(round((minutes % 1) * 60)):02d}"


# 擷取指定範圍的原始碼，移除共同縮排並保留原始行號。
# Extract a range of source lines, removing the common indentation and keeping the original line numbers.
def code_lines(file: str, start: int, end: int) -> list[tuple[int, str]]:
    lines = (ROOT / file).read_text(encoding="utf-8").splitlines()
    chunk = lines[start - 1 : end]
    indent = min(len(l) - len(l.lstrip()) for l in chunk if l.strip())
    return [(start + i, l[indent:]) for i, l in enumerate(chunk)]


# 複製討論提示卡所需圖片，並以 macOS sips 等比例縮至網頁使用寬度。
# Copy the figures needed for the cue cards and shrink them proportionally to page width with macOS sips.
def prepare_figures():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    for key, (source, _) in FIGURES.items():
        target = FIG_DIR / f"{key}.png"
        shutil.copyfile(ROOT / source, target)
        # sips 為 macOS 內建；只縮小寬度，保持比例。
        # sips is built into macOS; only shrink the width, keeping the aspect ratio.
        subprocess.run(
            ["sips", "--resampleWidth", str(WEB_WIDTH), str(target)],
            check=True,
            capture_output=True,
        )


# ---------------------------------------------------------------- 網頁 / Web page --------------
CSS = """
:root{
  /* Layout: sticky talk clock + time ribbon on top; each segment = figure column + cue column, code underneath. */
  --bg:#f4f6f8; --panel:#ffffff; --ink:#15212d; --muted:#5a6876; --line:#d9e0e6;
  --sens:#13748a; --res:#e4683f; --res-soft:#fbe7df; --sens-soft:#e1f0f3; --code-bg:#eef2f5; --mark:#fff1c7;
  --display:"IBM Plex Sans Condensed","Noto Sans TC","PingFang TC",system-ui,sans-serif;
  --body:"IBM Plex Sans","Noto Sans TC","PingFang TC",system-ui,sans-serif;
  --mono:"IBM Plex Mono",ui-monospace,Menlo,monospace;
}
@media (prefers-color-scheme: dark){:root:not([data-theme="light"]){
  --bg:#0f161d; --panel:#16202a; --ink:#e3e9ee; --muted:#93a2b0; --line:#2a3946;
  --sens:#4fb3c8; --res:#f0875f; --res-soft:#3a2219; --sens-soft:#14303a; --code-bg:#1b2731; --mark:#4a3d12; color-scheme:dark}}
:root[data-theme="dark"]{
  --bg:#0f161d; --panel:#16202a; --ink:#e3e9ee; --muted:#93a2b0; --line:#2a3946;
  --sens:#4fb3c8; --res:#f0875f; --res-soft:#3a2219; --sens-soft:#14303a; --code-bg:#1b2731; --mark:#4a3d12; color-scheme:dark}
*{box-sizing:border-box}
body{background:var(--bg);color:var(--ink);font:15px/1.6 var(--body);margin:0}
.wrap{max-width:1180px;margin:0 auto;padding-left:max(16px,env(safe-area-inset-left,0px));padding-right:max(16px,env(safe-area-inset-right,0px));padding-block:0 64px}
header.top{padding-block:36px 20px;display:grid;gap:10px}
.eyebrow{font:600 12px/1 var(--mono);letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}
h1{font:700 clamp(28px,4vw,40px)/1.15 var(--display);margin:0;text-wrap:balance}
.thesis{font:600 clamp(18px,2.4vw,22px)/1.4 var(--display);margin:0;color:var(--res);text-wrap:balance}
.thesis-zh{color:var(--muted);margin:0}
.howto{color:var(--muted);font-size:14px;margin:0;max-width:72ch}
.clock{position:sticky;top:env(safe-area-inset-top,0px);z-index:5;background:var(--bg);border-bottom:1px solid var(--line);padding-block:10px}
.clock-row{display:flex;flex-wrap:wrap;align-items:center;gap:12px}
.clock button{font:600 13px var(--body);border:1px solid var(--line);background:var(--panel);color:var(--ink);border-radius:6px;padding:6px 12px;cursor:pointer}
.clock button:focus-visible{outline:2px solid var(--sens);outline-offset:2px}
#elapsed{font:600 20px var(--mono);font-variant-numeric:tabular-nums;min-width:5ch}
#now{font-size:13px;color:var(--muted)}
.ribbon{display:flex;height:30px;margin-top:8px;border-radius:4px;overflow:hidden;position:relative}
.ribbon a{display:flex;align-items:center;justify-content:center;font:600 11px var(--body);color:var(--ink);text-decoration:none;background:var(--sens-soft);border-right:2px solid var(--bg);white-space:nowrap;overflow:hidden;min-width:0;padding:0 4px}
.ribbon a.active{background:var(--res);color:#fff}
#needle{position:absolute;top:0;bottom:0;width:2px;background:var(--ink);left:0;display:none}
section.seg{border-top:2px solid var(--ink);margin-top:40px;padding-top:14px;scroll-margin-top:110px}
.seg-head{display:flex;flex-wrap:wrap;align-items:baseline;gap:6px 16px}
.seg-head .t{font:600 14px var(--mono);color:var(--res);font-variant-numeric:tabular-nums}
.seg-head h2{font:700 24px/1.25 var(--display);margin:0;text-wrap:balance}
.goal{color:var(--muted);margin:6px 0 0;max-width:80ch}
.beat{display:grid;grid-template-columns:minmax(0,1.25fr) minmax(0,1fr);gap:20px 28px;margin-top:22px;align-items:start}
.beat.nofig{grid-template-columns:minmax(0,1fr)}
.beat.codebeat{grid-template-columns:minmax(0,1.1fr) minmax(0,1fr)}
@media (max-width:860px){.beat,.beat.codebeat{grid-template-columns:minmax(0,1fr)}}
figure{margin:0;background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:8px}
figure img{display:block;width:100%;height:auto;border-radius:3px;background:#fff}
figcaption{font-size:13px;color:var(--muted);padding:6px 2px 0}
.tag{display:inline-block;font:600 11px var(--mono);letter-spacing:.05em;text-transform:uppercase;padding:2px 6px;border-radius:3px;background:var(--code-bg);color:var(--muted);margin-left:6px}
.cues{display:grid;gap:14px;min-width:0}
.lbl{font:600 11px/1 var(--mono);letter-spacing:.08em;text-transform:uppercase;color:var(--muted);margin:0 0 6px}
ol.point{margin:0;padding-left:20px;font-size:14px}
ol.point li{margin:2px 0}
ul.says{list-style:none;margin:0;padding:0;display:grid;gap:8px}
ul.says li{border-left:3px solid var(--res);padding:2px 0 2px 10px}
.en{font-weight:600;display:block}
.zh{font-size:13px;color:var(--muted);display:block}
.chips{display:flex;flex-wrap:wrap;gap:6px}
.chip{font-size:13px;padding:2px 8px;border-radius:999px;border:1px solid var(--line);background:var(--panel)}
.num{font:500 13px/1.5 var(--mono);font-variant-numeric:tabular-nums;padding:2px 8px;border-radius:4px;background:var(--res-soft)}
.code{background:var(--code-bg);border-radius:6px;overflow-x:auto;min-width:0}
.code .ct{font-size:13px;font-weight:600;padding:8px 12px 0}
.code .cf{font:500 12px var(--mono);color:var(--muted);padding:0 12px}
pre{margin:0;padding:8px 0 10px;font:12.5px/1.55 var(--mono)}
pre span.l{display:block;padding:0 12px;white-space:pre}
pre span.l.hl{background:var(--mark)}
pre i{font-style:normal;color:var(--muted);display:inline-block;width:3.5ch;user-select:none}
.transition{margin-top:22px;background:var(--panel);border:1px dashed var(--line);border-radius:6px;padding:10px 14px}
.transition .en::before{content:"→ ";color:var(--res)}
.grid2{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,340px),1fr));gap:16px;margin-top:16px}
.qa{background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:12px 14px}
.qa .q{font-weight:600}
.qa .a{color:var(--muted);font-size:14px;margin-top:4px}
table.cheat{width:100%;border-collapse:collapse;font-size:14px;margin-top:12px}
table.cheat td{border-bottom:1px solid var(--line);padding:6px 8px;vertical-align:top}
table.cheat td:last-child{font-family:var(--mono);font-variant-numeric:tabular-nums;text-align:right;white-space:nowrap}
.optional{font:600 11px var(--mono);color:var(--sens);border:1px solid var(--sens);border-radius:3px;padding:1px 6px;margin-left:8px;vertical-align:middle}
.layers{display:flex;flex-wrap:wrap;align-items:stretch;gap:6px;margin-top:16px}
.layer{flex:1 1 150px;background:var(--panel);border:1px solid var(--line);border-top:3px solid var(--sens);border-radius:6px;padding:8px 10px;min-width:0}
.layer .ln{font:600 11px var(--mono);letter-spacing:.06em;color:var(--muted)}
.layer .lm{font:600 14px var(--mono);overflow-wrap:anywhere}
.layer .lr{font-size:13px;color:var(--muted)}
.layer-arrow{align-self:center;color:var(--res);font-weight:700}
@media (max-width:860px){.layer-arrow{display:none}}
.classes{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,360px),1fr));gap:16px;margin-top:18px}
.cls{background:var(--panel);border:1px solid var(--line);border-radius:6px;overflow:hidden;min-width:0}
.cls-h{background:var(--sens-soft);padding:8px 12px;border-bottom:1px solid var(--line)}
.cls-h .k{font:500 11px var(--mono);color:var(--muted)}
.cls-h .n{font:700 17px var(--mono)}
.cls-h .f{font:500 11px var(--mono);color:var(--muted);float:right}
.cls-role{font-size:13px;color:var(--muted);padding:8px 12px 2px}
.cls-part{border-top:1px dashed var(--line);margin-top:6px;padding:6px 12px 8px}
.cls-part .lbl{margin-bottom:4px}
.mem{padding:3px 0}
.mem code{font:500 12.5px/1.45 var(--mono);display:block;overflow-wrap:anywhere}
.mem span{font-size:13px;color:var(--muted);display:block}
ol.flow{margin:12px 0 0;padding-left:22px;display:grid;gap:6px}
ol.flow code{font:600 13px var(--mono);background:var(--code-bg);padding:1px 6px;border-radius:3px;overflow-wrap:anywhere}
ol.flow span{display:block;font-size:14px;color:var(--muted)}
h3.sub{font:700 17px/1.3 var(--display);margin:26px 0 0}
@media (prefers-reduced-motion:reduce){*{scroll-behavior:auto!important}}
"""

JS = """
const segs = %SEGS%;
const total = 20*60;
let t0 = null, acc = 0, timer = null;
const el = id => document.getElementById(id);
function fmt(s){s=Math.max(0,Math.floor(s));return Math.floor(s/60)+':'+String(s%60).padStart(2,'0');}
function tick(){
  const s = acc + (t0 ? (Date.now()-t0)/1000 : 0);
  el('elapsed').textContent = fmt(s);
  const m = s/60;
  const cur = segs.find(x => m >= x.start && m < x.end) || segs[segs.length-1];
  document.querySelectorAll('.ribbon a').forEach(a => a.classList.toggle('active', a.dataset.id===cur.id));
  const left = cur.end*60 - s;
  el('now').textContent = s>=total ? '已超過 20 分鐘' : '現在：'+cur.title+'（本段剩 '+fmt(left)+'）';
  const n = el('needle'); n.style.display='block'; n.style.left = Math.min(100, s/total*100)+'%';
}
el('go').addEventListener('click', () => {
  if (t0){acc += (Date.now()-t0)/1000; t0=null; clearInterval(timer); el('go').textContent='繼續';}
  else {t0=Date.now(); timer=setInterval(tick,500); el('go').textContent='暫停'; tick();}
});
el('reset').addEventListener('click', () => {
  acc=0; t0=null; clearInterval(timer); el('go').textContent='開始計時';
  el('elapsed').textContent='0:00'; el('now').textContent='按「開始計時」後，會標出目前該講的段落'; el('needle').style.display='none';
  document.querySelectorAll('.ribbon a').forEach(a => a.classList.remove('active'));
});
"""


# 跳脫文字中的 HTML 特殊字元及引號，供網頁內容安全插入。
# Escape HTML special characters and quotes so text can be inserted into the page safely.
def esc(text: str) -> str:
    return html.escape(text, quote=True)


# 將指定原始碼片段轉為含行號與重點標記的 HTML 區塊。
# Turn a source excerpt into an HTML block with line numbers and highlighted lines.
def html_code(code, title) -> str:
    file, start, end, hl = code
    rows = "".join(
        f'<span class="l{" hl" if n in hl else ""}"><i>{n}</i>{esc(line)}</span>'
        for n, line in code_lines(file, start, end)
    )
    return (
        f'<div class="code"><div class="ct">{esc(title)}</div><div'
        f' class="cf">{esc(file)}:{start}–{end}　黃底 ='
        f" 要指的行</div><pre>{rows}</pre></div>"
    )


# 將指圖順序、關鍵句與補充提示轉為 HTML 提示區塊。
# Turn the pointing order, key sentences and extra cues into an HTML cue block.
def html_cues(beat) -> str:
    parts = []
    if beat.get("point"):
        items = "".join(f"<li>{esc(p)}</li>" for p in beat["point"])
        parts.append(
            f'<div><p class="lbl">指圖順序</p><ol class="point">{items}</ol></div>'
        )
    says = "".join(
        f'<li><span class="en">{esc(en)}</span><span class="zh">{esc(zh)}</span></li>'
        for en, zh in beat["says"]
    )
    parts.append(f'<div><p class="lbl">關鍵句</p><ul class="says">{says}</ul></div>')
    if beat.get("concepts"):
        chips = "".join(f'<span class="chip">{esc(c)}</span>' for c in beat["concepts"])
        parts.append(
            f'<div><p class="lbl">概念錨點</p><div class="chips">{chips}</div></div>'
        )
    if beat.get("numbers"):
        nums = "".join(f'<span class="num">{esc(n)}</span>' for n in beat["numbers"])
        parts.append(
            f'<div><p class="lbl">要引用的數字</p><div class="chips">{nums}</div></div>'
        )
    return f'<div class="cues">{"".join(parts)}</div>'


# 將單一討論節拍的圖片、提示與程式碼組成 HTML 內容。
# Combine one beat's figure, cues and code into HTML.
def html_beat(beat) -> str:
    cues = html_cues(beat)
    if beat.get("fig"):
        _, caption = FIGURES[beat["fig"]]
        opt = '<span class="optional">時間夠才講</span>' if beat.get("optional") else ""
        fig = (
            f'<figure><img src="figs/{beat["fig"]}.png" alt="{esc(caption)}"'
            f' loading="lazy"><figcaption>{esc(caption)}{opt}</figcaption></figure>'
        )
        return f'<div class="beat">{fig}{cues}</div>'
    if beat.get("code"):
        return (
            '<div class="beat'
            f' codebeat">{html_code(beat["code"], beat["code_title"])}{cues}</div>'
        )
    return f'<div class="beat nofig">{cues}</div>'


# 建立包含分層架構、模組職責與函式索引的 HTML 程式導覽。
# HTML code tour with the layered architecture, module responsibilities and function index.
def html_arch() -> str:
    layers = '<span class="layer-arrow">→</span>'.join(
        f'<div class="layer"><div class="ln">{esc(n)}</div><div'
        f' class="lm">{esc(m)}</div><div class="lr">{esc(r)}</div></div>'
        for n, m, r in ARCH_LAYERS
    )

    # 將名稱與說明配對轉為帶標題的 HTML 清單，空清單則回傳空字串。
    # Turn (name, description) pairs into a titled HTML list; empty input returns an empty string.
    def members(label, rows):
        if not rows:
            return ""
        items = "".join(
            f'<div class="mem"><code>{esc(a)}</code><span>{esc(b)}</span></div>'
            for a, b in rows
        )
        return f'<div class="cls-part"><p class="lbl">{label}</p>{items}</div>'

    cards = "".join(
        f'<div class="cls"><div class="cls-h"><span class="f">{esc(c["file"])}</span>'
        f'<div class="k">{esc(c["kind"])}</div><div class="n">{esc(c["name"])}</div></div>'
        f'<div class="cls-role">{esc(c["role"])}</div>'
        f'{members("屬性", c["attrs"])}{members("方法 → 回傳", c["methods"])}</div>'
        for c in ARCH_CLASSES
    )
    flow = "".join(
        f"<li><code>{esc(a)}</code><span>{esc(b)}</span></li>" for a, b in ARCH_FLOW
    )
    says = "".join(
        f'<li><span class="en">{esc(en)}</span><span class="zh">{esc(zh)}</span></li>'
        for en, zh in ARCH_SAYS
    )
    return (
        '<section class="seg" id="arch"><div class="seg-head"><span class="t">講前先讀</span>'
        "<h2>程式導覽：架構與函式</h2></div>"
        f'<p class="goal">{esc(ARCH_INTRO)}</p><div class="layers">{layers}</div>'
        f'<h3 class="sub">類別卡</h3><div class="classes">{cards}</div>'
        f'<h3 class="sub">執行一次 main.py 發生什麼事</h3><ol class="flow">{flow}</ol>'
        '<h3 class="sub">被問到程式架構時</h3>'
        f'<ul class="says" style="margin-top:10px">{says}</ul></section>'
    )


# 組合時間軸、程式導覽、討論段落、問答與計時功能，回傳完整提示卡 HTML。
# Assemble the timeline, code tour, sections, Q&A and timer into the complete cue-card HTML.
def build_html() -> str:
    ribbon = "".join(
        f'<a href="#{s["id"]}" data-id="{s["id"]}" style="flex:{s["end"] - s["start"]}"'
        f' title="{esc(s["title"])}">{fmt_time(s["start"])}</a>'
        for s in SEGMENTS
    )
    sections = []
    for s in SEGMENTS:
        beats = "".join(html_beat(b) for b in s["beats"])
        en, zh = s["transition"]
        sections.append(
            f'<section class="seg" id="{s["id"]}"><div class="seg-head"><span'
            f' class="t">{fmt_time(s["start"])}–{fmt_time(s["end"])}・{s["end"] - s["start"]:g}'
            f' 分</span><h2>{esc(s["title"])}</h2></div><p'
            f' class="goal">{esc(s["goal"])}</p>{beats}<div class="transition"><p'
            f' class="lbl">銜接句</p><span class="en">{esc(en)}</span><span'
            f' class="zh">{esc(zh)}</span></div></section>'
        )
    qa = "".join(
        f'<div class="qa"><div class="q">{esc(q)}</div><div'
        f' class="a">{esc(a)}</div></div>'
        for q, a in QA
    )
    cheat = "".join(f"<tr><td>{esc(k)}</td><td>{esc(v)}</td></tr>" for k, v in CHEAT)
    segs_js = str(
        [
            dict(id=s["id"], start=s["start"], end=s["end"], title=s["title"])
            for s in SEGMENTS
        ]
    ).replace("'", '"')
    return f"""<title>腫瘤模擬討論提示卡</title>
<meta name="claude:viewport-fit-sides" content="cover">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;500;600&family=IBM+Plex+Sans+Condensed:wght@600;700&family=IBM+Plex+Sans:wght@400;600&family=Noto+Sans+TC:wght@400;600;700&display=swap">
<style>{CSS}</style>
<div class="wrap">
<header class="top">
  <div class="eyebrow">BME 6120 · Quantitative tumor modeling · 20-minute discussion</div>
  <h1>腫瘤分枝模型：與老師討論的提示卡</h1>
  <p class="thesis">{esc(THESIS[0])}</p>
  <p class="thesis-zh">{esc(THESIS[1])}</p>
  <p class="howto">用法：粗體英文是可以直接說出口的關鍵句，灰字是意思提醒；「指圖順序」告訴你手指要先指哪裡；黃底程式行是要對照的重點。標「時間夠才講」的部分可以跳過。</p>
</header>
<div class="clock"><div class="clock-row">
  <button id="go" type="button">開始計時</button><button id="reset" type="button">歸零</button>
  <span id="elapsed">0:00</span><span id="now">按「開始計時」後，會標出目前該講的段落</span>
</div><div class="ribbon">{ribbon}<span id="needle"></span></div></div>
{html_arch()}
{"".join(sections)}
<section class="seg" id="qa"><div class="seg-head"><span class="t">備用</span><h2>老師可能問的問題</h2></div>
<p class="goal">先用一句英文回答核心，再指向對應的圖或程式行。</p><div class="grid2">{qa}</div></section>
<section class="seg" id="cheat"><div class="seg-head"><span class="t">備用</span><h2>數字小抄</h2></div>
<p class="goal">所有數字來自 outputs/ 內的 CSV／JSON（預設參數、每組 30 次試驗，圖 08 每點 200 次）。</p>
<table class="cheat">{cheat}</table></section>
</div>
<script>{JS.replace("%SEGS%", segs_js)}</script>
"""


# ---------------------------------------------------------------- Word ---------------
INK = RGBColor(0x15, 0x21, 0x2D)
MUTED = RGBColor(0x5A, 0x68, 0x76)
ACCENT = RGBColor(0xC9, 0x52, 0x2B)
TEAL = RGBColor(0x13, 0x74, 0x8A)


# 設定 Word 文字片段的西文與中文字型、字級、顏色及粗體。
# Set the Latin and East Asian fonts, size, colour and bold of a Word text run.
def set_fonts(
    run,
    latin="Arial",
    east="PingFang TC",
    size=None,
    bold=None,
    color=None,
    italic=None,
):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), latin)
    fonts.set(qn("w:eastAsia"), east)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


# 在 Word 段落屬性中加入指定底色。
# Add a shading colour to a Word paragraph's properties.
def shade(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


# 為 Word 段落加入指定顏色的左側框線。
# Add a coloured left border to a Word paragraph.
def left_border(paragraph, color="E4683F"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (
        ("w:val", "single"),
        ("w:sz", "18"),
        ("w:space", "6"),
        ("w:color", color),
    ):
        left.set(qn(k), v)
    borders.append(left)
    ppr.append(borders)


# 建立 Word 段落，設定段後距離並依選項加入格式化文字。
# Create a Word paragraph with spacing after and optional formatted text.
def para(doc, text="", size=10.5, color=INK, bold=False, space_after=4, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        set_fonts(p.add_run(text), size=size, color=color, bold=bold)
    return p


# 建立帶統一間距與字型的 Word 小標題段落。
# Create a Word subheading paragraph with the shared spacing and font.
def label(doc, text):
    p = para(doc, space_after=2)
    p.paragraph_format.space_before = Pt(6)
    set_fonts(p.add_run(text), size=8.5, color=MUTED, bold=True)
    return p


# 將討論節拍的指圖順序、關鍵句與補充提示寫入 Word 文件。
# Write a beat's pointing order, key sentences and extra cues into the Word document.
def word_cues(doc, beat):
    if beat.get("point"):
        label(doc, "指圖順序")
        for item in beat["point"]:
            p = para(doc, style="List Number", space_after=1)
            set_fonts(p.add_run(item), size=10)
    label(doc, "關鍵句")
    for en, zh in beat["says"]:
        p = para(doc, space_after=3)
        p.paragraph_format.left_indent = Cm(0.3)
        left_border(p)
        set_fonts(p.add_run(en), size=10.5, bold=True)
        p.add_run().add_break()
        set_fonts(p.add_run(zh), size=9.5, color=MUTED)
    if beat.get("concepts"):
        label(doc, "概念錨點")
        para(doc, "・".join(beat["concepts"]), size=10, color=TEAL)
    if beat.get("numbers"):
        label(doc, "要引用的數字")
        p = para(doc, space_after=4)
        set_fonts(
            p.add_run("　｜　".join(beat["numbers"])),
            latin="Menlo",
            size=9.5,
            color=ACCENT,
        )


# 將帶標題、原始行號與重點標記的程式碼片段寫入 Word 文件。
# Write a titled code excerpt with original line numbers and highlighted lines into the Word document.
def word_code(doc, code, title):
    file, start, end, hl = code
    p = para(doc, space_after=0)
    p.paragraph_format.space_before = Pt(8)
    set_fonts(p.add_run(title), size=10, bold=True)
    para(
        doc,
        f"{file}:{start}–{end}（粗體黃底 = 要指的行）",
        size=8.5,
        color=MUTED,
        space_after=2,
    )
    for n, line in code_lines(file, start, end):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing = 1.0
        shade(q, "FFF1C7" if n in hl else "EEF2F5")
        set_fonts(q.add_run(f"{n:>4}  "), latin="Menlo", size=8, color=MUTED)
        set_fonts(q.add_run(line or " "), latin="Menlo", size=8, bold=n in hl)


# 將名稱與說明配對寫成 Word 表格，空清單則略過。
# Write (name, description) pairs as a Word table; skip empty input.
def word_members(doc, label_text, rows):
    if not rows:
        return
    label(doc, label_text)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for a, b in rows:
        row = table.add_row().cells
        row[0].text = row[1].text = ""
        set_fonts(row[0].paragraphs[0].add_run(a), latin="Menlo", size=8.5)
        set_fonts(row[1].paragraphs[0].add_run(b), size=9, color=MUTED)
        row[0].width, row[1].width = Cm(8.6), Cm(9.2)


# 將分層架構、模組職責與函式索引寫入 Word 程式導覽。
# Write the layered architecture, module responsibilities and function index into the Word code tour.
def word_arch(doc):
    h = doc.add_heading(level=1)
    set_fonts(h.add_run("講前先讀　"), size=13, color=ACCENT)
    set_fonts(h.add_run("程式導覽：架構與函式"), size=16, color=INK)
    para(doc, ARCH_INTRO, size=10, color=MUTED, space_after=6)
    label(doc, "分層（由入口往下呼叫）")
    para(
        doc,
        "  →  ".join(m for _, m, _ in ARCH_LAYERS),
        size=10,
        color=TEAL,
        bold=True,
    )
    for n, m, r in ARCH_LAYERS:
        p = para(doc, space_after=1)
        set_fonts(p.add_run(f"{n}　{m}　"), latin="Menlo", size=9, bold=True)
        set_fonts(p.add_run(r), size=9, color=MUTED)
    for c in ARCH_CLASSES:
        p = para(doc, space_after=0)
        p.paragraph_format.space_before = Pt(12)
        shade(p, "E1F0F3")
        set_fonts(p.add_run(c["kind"] + "  "), latin="Menlo", size=8.5, color=MUTED)
        set_fonts(p.add_run(c["name"]), latin="Menlo", size=13, bold=True)
        set_fonts(p.add_run("    " + c["file"]), latin="Menlo", size=8.5, color=MUTED)
        para(doc, c["role"], size=9.5, color=MUTED, space_after=2)
        word_members(doc, "屬性", c["attrs"])
        word_members(doc, "方法 → 回傳", c["methods"])
    label(doc, "執行一次 main.py 發生什麼事")
    for a, b in ARCH_FLOW:
        p = para(doc, style="List Number", space_after=2)
        set_fonts(p.add_run(a), latin="Menlo", size=9, bold=True)
        p.add_run().add_break()
        set_fonts(p.add_run(b), size=9.5, color=MUTED)
    label(doc, "被問到程式架構時")
    for en, zh in ARCH_SAYS:
        p = para(doc, space_after=3)
        p.paragraph_format.left_indent = Cm(0.3)
        left_border(p)
        set_fonts(p.add_run(en), size=10.5, bold=True)
        p.add_run().add_break()
        set_fonts(p.add_run(zh), size=9.5, color=MUTED)


# 建立含版面樣式、程式導覽、討論提示與備用資料的 Word 文件並儲存。
# Build and save the Word document with page styles, code tour, cues and backup material.
def build_docx(path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # US Letter
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.9))
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang TC")

    p = para(
        doc,
        "BME 6120 · Quantitative tumor modeling · 20-minute discussion",
        size=9,
        color=MUTED,
    )
    p = para(doc, space_after=6)
    set_fonts(p.add_run("腫瘤分枝模型：與老師討論的提示卡"), size=22, bold=True)
    p = para(doc, space_after=0)
    set_fonts(p.add_run(THESIS[0]), size=13, bold=True, color=ACCENT)
    para(doc, THESIS[1], size=10.5, color=MUTED, space_after=8)
    para(
        doc,
        "用法：粗體英文是可直接說出口的關鍵句，"
        "灰字是意思提醒；「指圖順序」是手指要先指的地方；"
        "程式中黃底粗體行是要對照的重點；標「時間夠才講」可跳過。",
        size=9.5,
        color=MUTED,
        space_after=8,
    )

    # 時間表
    # Timetable
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("時間", "段落", "打開的圖")):
        cell.text = ""
        set_fonts(cell.paragraphs[0].add_run(text), size=9.5, bold=True)
    for s in SEGMENTS:
        figs = [
            " ".join(FIGURES[b["fig"]][1].split(" ")[:2])
            + (" (選)" if b.get("optional") else "")
            for b in s["beats"]
            if b.get("fig")
        ]
        row = table.add_row().cells
        values = (
            f"{fmt_time(s['start'])}–{fmt_time(s['end'])}",
            s["title"],
            "、".join(figs) or "—",
        )
        for cell, text in zip(row, values):
            cell.text = ""
            set_fonts(cell.paragraphs[0].add_run(text), size=9.5)
    for row in table.rows:
        for cell, width in zip(row.cells, (Cm(2.6), Cm(7.6), Cm(7.6))):
            cell.width = width

    doc.add_page_break()
    word_arch(doc)

    for s in SEGMENTS:
        doc.add_page_break()
        h = doc.add_heading(level=1)
        set_fonts(
            h.add_run(f"{fmt_time(s['start'])}–{fmt_time(s['end'])}　"),
            size=13,
            color=ACCENT,
        )
        set_fonts(h.add_run(s["title"]), size=16, color=INK)
        para(doc, s["goal"], size=10, color=MUTED, space_after=6)
        for beat in s["beats"]:
            if beat.get("fig"):
                source, caption = FIGURES[beat["fig"]]
                p = para(doc, space_after=2)
                p.paragraph_format.space_before = Pt(10)
                set_fonts(p.add_run(caption), size=11, bold=True, color=TEAL)
                if beat.get("optional"):
                    set_fonts(p.add_run("　時間夠才講"), size=9, color=MUTED)
                pic = doc.add_paragraph()
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(
                    str(FIG_DIR / f"{beat['fig']}.png"), width=Cm(17.5)
                )
            if beat.get("code"):
                word_code(doc, beat["code"], beat["code_title"])
            word_cues(doc, beat)
        en, zh = s["transition"]
        label(doc, "銜接句")
        p = para(doc, space_after=0)
        shade(p, "F4F6F8")
        set_fonts(p.add_run("→ " + en), size=10.5, bold=True)
        p.add_run().add_break()
        set_fonts(p.add_run(zh), size=9.5, color=MUTED)

    doc.add_page_break()
    h = doc.add_heading(level=1)
    set_fonts(h.add_run("老師可能問的問題"), size=16, color=INK)
    for q, a in QA:
        p = para(doc, space_after=0)
        p.paragraph_format.space_before = Pt(6)
        set_fonts(p.add_run("Q  " + q), size=10.5, bold=True)
        para(doc, a, size=10, color=MUTED, space_after=2)

    h = doc.add_heading(level=1)
    set_fonts(h.add_run("數字小抄"), size=16, color=INK)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for k, v in CHEAT:
        row = table.add_row().cells
        row[0].text = row[1].text = ""
        set_fonts(row[0].paragraphs[0].add_run(k), size=10)
        set_fonts(row[1].paragraphs[0].add_run(v), latin="Menlo", size=10, color=ACCENT)
        row[0].width, row[1].width = Cm(9), Cm(8.8)
    para(
        doc,
        "數字來自 outputs/ 內的 CSV／JSON（預設參數、每組 30 次試驗；圖 08 每點"
        " 200 次）。",
        size=8.5,
        color=MUTED,
    )
    doc.save(path)


# 準備圖片並輸出討論提示卡的 HTML 與 Word 檔案。
# Prepare the figures and write the cue cards as HTML and Word.
def main():
    OUT.mkdir(parents=True, exist_ok=True)
    prepare_figures()
    (OUT / "discussion_guide.html").write_text(build_html(), encoding="utf-8")
    build_docx(OUT / "腫瘤模擬_20分鐘討論提示卡.docx")
    print("wrote", OUT)


if __name__ == "__main__":
    main()
